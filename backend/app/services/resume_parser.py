import logging

import fitz
import docx

logger = logging.getLogger(__name__)


def parse_resume(file_path: str) -> str:
    """Parse a PDF or DOCX resume file and return its text content."""
    try:
        if file_path.lower().endswith(".docx"):
            return _parse_docx(file_path)
        return _parse_pdf(file_path)
    except ValueError:
        raise
    except Exception as e:
        logger.error("Failed to parse resume '%s': %s", file_path, e, exc_info=True)
        raise ValueError(f"Failed to parse resume: {e}") from e


def _parse_docx(file_path: str) -> str:
    try:
        doc = docx.Document(file_path)
    except Exception as e:
        raise ValueError(f"Cannot open DOCX file: {e}") from e

    text = []
    for para in doc.paragraphs:
        if para.text.strip():
            text.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text.append(cell.text.strip())
    return "\n".join(text)


def _parse_pdf(file_path: str) -> str:
    try:
        document = fitz.open(file_path)
    except Exception as e:
        raise ValueError(f"Cannot open PDF file: {e}") from e

    try:
        all_text = []
        for page in document:
            page_text = _extract_page_columns(page)
            if page_text.strip():
                all_text.append(page_text)
        return "\n".join(all_text)
    finally:
        document.close()


def _extract_page_columns(page) -> str:
    """Extract text from a PDF page, handling multi-column layouts."""
    data = page.get_text("dict", sort=True)
    blocks = data.get("blocks", [])

    # Get text blocks with bounding boxes
    text_blocks = []
    for block in blocks:
        if block.get("type") != 0:  # skip images
            continue
        lines_text = []
        for line in block.get("lines", []):
            spans_text = "".join(span.get("text", "") for span in line.get("spans", []))
            if spans_text.strip():
                lines_text.append(spans_text.strip())
        if lines_text:
            text_blocks.append({
                "x0": block["bbox"][0],
                "y0": block["bbox"][1],
                "x1": block["bbox"][2],
                "y1": block["bbox"][3],
                "text": "\n".join(lines_text),
            })

    if not text_blocks:
        return ""

    # Detect if page has multiple columns
    page_width = page.rect.width
    mid_xs = sorted([(b["x0"] + b["x1"]) / 2 for b in text_blocks])

    left_blocks = []
    right_blocks = []

    # Find column boundary
    column_split = _find_column_split(text_blocks, page_width)

    if column_split is not None:
        for b in text_blocks:
            block_mid = (b["x0"] + b["x1"]) / 2
            if block_mid < column_split:
                left_blocks.append(b)
            else:
                right_blocks.append(b)

        # Sort each column top-to-bottom
        left_blocks.sort(key=lambda b: b["y0"])
        right_blocks.sort(key=lambda b: b["y0"])

        left_text = "\n".join(b["text"] for b in left_blocks)
        right_text = "\n".join(b["text"] for b in right_blocks)
        return left_text + "\n" + right_text
    else:
        # Single column: sort top-to-bottom
        text_blocks.sort(key=lambda b: (b["y0"], b["x0"]))
        return "\n".join(b["text"] for b in text_blocks)


def _find_column_split(text_blocks, page_width):
    """Find the x-coordinate that splits two columns, or None if single-column."""
    if len(text_blocks) < 4:
        return None

    # Only look in the middle 60% of the page
    min_split = page_width * 0.20
    max_split = page_width * 0.80
    min_gap = page_width * 0.03

    # Block horizontal extents
    intervals = [(b["x0"], b["x1"]) for b in text_blocks]

    # Find gaps between block edges
    edges = set()
    for x0, x1 in intervals:
        edges.add(x0)
        edges.add(x1)
    edges = sorted(edges)

    best_split = None
    best_gap = 0

    for i in range(len(edges) - 1):
        gap_start = edges[i]
        gap_end = edges[i + 1]
        gap_mid = (gap_start + gap_end) / 2
        gap_width = gap_end - gap_start

        if gap_mid < min_split or gap_mid > max_split:
            continue
        if gap_width < min_gap:
            continue

        # Ensure no block crosses this gap
        spans_gap = any(x0 < gap_mid < x1 for x0, x1 in intervals)
        if spans_gap:
            continue

        # Count blocks on each side
        left_count = sum(1 for x0, x1 in intervals if (x0 + x1) / 2 < gap_mid)
        right_count = sum(1 for x0, x1 in intervals if (x0 + x1) / 2 >= gap_mid)

        # Both sides need content
        if left_count >= 2 and right_count >= 2 and gap_width > best_gap:
            best_gap = gap_width
            best_split = gap_mid

    return best_split